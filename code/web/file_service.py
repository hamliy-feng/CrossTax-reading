"""Text extraction with original bytes and precise page/sheet/line locators."""
from __future__ import annotations
import csv
import io
from pathlib import Path
import zipfile

EXTENSIONS={'.pdf','.docx','.txt','.md','.csv','.xlsx'}
MAX_BYTES=25*1024*1024


def parse(name,body):
    suffix=Path(name).suffix.lower()
    if suffix not in EXTENSIONS:
        raise ValueError('支持 PDF、DOCX、TXT、MD、CSV、XLSX')
    if len(body)>MAX_BYTES:
        raise ValueError('单文件不能超过 25 MiB')
    if suffix in {'.docx','.xlsx'}:
        with zipfile.ZipFile(io.BytesIO(body)) as z:
            if sum(i.file_size for i in z.infolist())>80*1024*1024:
                raise ValueError('压缩文档展开后过大')
    chunks=[];partial=False
    if suffix=='.pdf':
        from pypdf import PdfReader
        reader=PdfReader(io.BytesIO(body))
        if reader.is_encrypted:
            raise ValueError('请上传未加密的 PDF')
        for n,p in enumerate(reader.pages[:500],1):
            text=p.extract_text() or ''
            if text.strip():
                chunks.append({'locator':f'page:{n}','page':n,'text':text})
        partial=len(reader.pages)>500
    elif suffix=='.docx':
        from docx import Document
        d=Document(io.BytesIO(body))
        chunks=[{'locator':f'paragraph:{i}','text':p.text} for i,p in enumerate(d.paragraphs,1) if p.text.strip()]
        for i,t in enumerate(d.tables,1):
            for j,row in enumerate(t.rows,1):
                chunks.append({'locator':f'table:{i}/row:{j}','text':' | '.join(c.text for c in row.cells)})
    elif suffix=='.xlsx':
        from openpyxl import load_workbook
        book=load_workbook(io.BytesIO(body),read_only=True,data_only=False,keep_links=False)
        try:
            for sheet in book:
                for n,row in enumerate(sheet.iter_rows(values_only=True),1):
                    if n>20000:
                        partial=True;break
                    if any(v is not None for v in row):
                        chunks.append({'locator':f'sheet:{sheet.title}/row:{n}','sheet':sheet.title,'row':n,'text':' | '.join('' if v is None else str(v) for v in row[:200])})
        finally:
            book.close()
    else:
        try:
            text=body.decode('utf-8-sig')
        except UnicodeDecodeError:
            text=body.decode('gb18030')
        if suffix=='.csv':
            for n,row in enumerate(csv.reader(io.StringIO(text)),1):
                if n>20000:
                    partial=True;break
                chunks.append({'locator':f'row:{n}','row':n,'text':' | '.join(row)})
        else:
            lines=text.splitlines()
            chunks=[{'locator':f'lines:{i+1}-{min(i+100,len(lines))}','text':'\n'.join(lines[i:i+100])} for i in range(0,len(lines),100)]
    kept=[];size=0
    for x in chunks:
        size+=len(x['text'])
        if size>1500000:
            partial=True;break
        kept.append(x)
    return {'chunks':kept,'status':'partial' if partial else 'parsed' if kept else 'needs_text',
        'note':'解析容量上限外的内容保留在原件中，可拆分再上传' if partial else '未提取文字；扫描件需文字识别后再使用' if not kept else '',
        'parser':'text extraction; no external OCR','text_characters':sum(len(x['text']) for x in kept)}
