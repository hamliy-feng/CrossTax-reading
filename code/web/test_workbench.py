"""Behavioral regression: authorization, persistence, branches and real tools.

Model/auth network dependencies are mocked here; separate acceptance runs record
actual providers. The optional local legal query uses the dedicated read-only DB.
"""
import asyncio
import io
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch, AsyncMock
from fastapi.testclient import TestClient
import workbench as w
from workspace_store import Store, Blobs, Conflict
import legal_service


async def fake_stream(messages,config):
    yield {'type':'delta','text':'已保存的回答。'}
    yield {'type':'delta','text':w.model_client.DISCLAIMER}


class WorkbenchTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory()
        self.db=Store(path=Path(self.tmp.name)/'app.sqlite3',dsn='')
        self.blob=Blobs();self.blob.client=None;self.blob.local=Path(self.tmp.name)/'files'
        self.patches=[patch.object(w,'store',self.db),patch.object(w,'blobs',self.blob),
            patch.object(w.server,'_resolved_config',return_value=('deepseek','test-model','test-key')),
            patch.object(w.model_client,'stream',fake_stream)]
        for p in self.patches:p.start()
        w.running.clear();w.requests_by_owner.clear()
        self.client=TestClient(w.app);self.other=TestClient(w.app)

    def tearDown(self):
        self.client.close();self.other.close()
        for p in reversed(self.patches):p.stop()
        self.tmp.cleanup()

    def case(self,id=1,title='项目对话',client=None):
        r=(client or self.client).post('/api/cases',json={'id':id,'title':title,'facts':{}})
        self.assertEqual(r.status_code,200,r.text);return r.json()['case']

    def chat(self,id=1,**kwargs):
        r=self.client.post('/api/research',json={'case_id':id,'prompt':'你好',**kwargs})
        self.assertEqual(r.status_code,200,r.text)
        events=[json.loads(x) for x in r.text.splitlines()]
        self.assertEqual(events[-1]['type'],'finish',events)
        return events[-1]['case']

    def test_owner_isolation_all_records(self):
        self.case();self.assertEqual(self.other.get('/api/cases/1').status_code,404)
        f=self.client.post('/api/files',files={'file':('facts.md',b'private text')}).json()
        self.assertEqual(self.other.get('/api/files/'+f['id']+'/download').status_code,404)
        self.assertEqual(self.other.delete('/api/files/'+f['id']).status_code,404)

    def test_revision_conflict_does_not_overwrite(self):
        c=self.case();r=self.client.post('/api/cases',json={**c,'title':'新标题'});self.assertEqual(r.status_code,200)
        r=self.client.post('/api/cases',json={**c,'title':'过时修改'});self.assertEqual(r.status_code,409)
        self.assertEqual(self.client.get('/api/cases/1').json()['title'],'新标题')

    def test_chat_saved_by_server_and_disclaimer_once(self):
        self.case();c=self.chat();self.assertEqual(len(c['messages']),2)
        self.assertEqual(c['messages'][-1]['text'].count(w.model_client.DISCLAIMER),1)
        r=self.client.post('/api/cases',json={**c,'messages':[]});self.assertEqual(len(r.json()['case']['messages']),2)
        self.assertEqual(self.client.get('/api/cases/1/export').text.count(w.model_client.DISCLAIMER),1)

    def test_edit_regenerate_preserve_branches_and_context(self):
        self.case();c=self.chat();u,a=c['messages'];original=a['id']
        c=self.chat(operation='regenerate',message_id=original)
        self.assertEqual(len(c['messages']),3);self.assertEqual(c['messages'][-1]['parent_id'],u['id'])
        c=self.chat(prompt='换一个问题',operation='edit',message_id=u['id'])
        self.assertEqual(len(c['messages']),5);self.assertIsNone(c['messages'][-2]['parent_id'])
        self.assertEqual([x['text'] for x in w.active_messages(c) if x['role']=='user'],['换一个问题'])
        c=self.client.post('/api/cases/1/branch',json={'message_id':original}).json()
        self.assertEqual(w.active_messages(c)[0]['text'],'你好')

    def test_delete_message_returns_context_to_parent(self):
        self.case();c=self.chat();u=c['messages'][0]
        c=self.client.delete('/api/cases/1/messages/'+u['id']).json()
        self.assertEqual(w.active_messages(c),[]);self.assertEqual(len(c['messages']),2)

    def test_project_moves_atomic_and_delete_keeps_chats(self):
        self.case();self.case(2);p=self.client.post('/api/projects',json={'title':'方案'}).json()
        r=self.client.post('/api/projects/move',json={'case_ids':[1,99],'project_id':p['id']});self.assertEqual(r.status_code,400)
        self.assertIsNone(self.client.get('/api/cases/1').json()['project_id'])
        self.assertEqual(self.client.post('/api/projects/move',json={'case_ids':[1,2],'project_id':p['id']}).status_code,200)
        self.assertEqual(self.other.patch('/api/projects/'+p['id'],json={'title':'入侵'}).status_code,404)
        self.client.delete('/api/projects/'+p['id']);rows=self.client.get('/api/cases').json()['cases']
        self.assertEqual(len(rows),2);self.assertTrue(all(c['project_id'] is None for c in rows))

    def test_file_formats_locators_and_originals(self):
        from docx import Document
        from openpyxl import Workbook
        from pypdf import PdfWriter
        d=Document();d.add_paragraph('合同段落');db=io.BytesIO();d.save(db)
        b=Workbook();b.active.append(['收入','123.45']);xb=io.BytesIO();b.save(xb);b.close()
        p=PdfWriter();p.add_blank_page(100,100);pb=io.BytesIO();p.write(pb)
        fixtures=[('a.txt','第一行\n第二行'.encode()),('a.md',b'# title'),('a.csv',b'name,value\nx,1'),('a.docx',db.getvalue()),('a.xlsx',xb.getvalue()),('a.pdf',pb.getvalue())]
        for name,body in fixtures:
            r=self.client.post('/api/files',files={'file':(name,body)});self.assertEqual(r.status_code,200,r.text)
            f=r.json();self.assertEqual(self.client.get('/api/files/'+f['id']+'/download').content,body)
            if name.endswith('pdf'):self.assertEqual(f['status'],'needs_text')
            else:self.assertTrue(f['chunks'][0]['locator'])

    def test_attachment_size_limit_and_unknown_format(self):
        self.assertEqual(self.client.post('/api/files',files={'file':('a.exe',b'bad')}).status_code,400)
        r=self.client.post('/api/files',files={'file':('a.txt',b'a'*(25*1024*1024+1))});self.assertEqual(r.status_code,413)

    def test_parse_failure_keeps_original(self):
        r=self.client.post('/api/files',files={'file':('broken.docx',b'corrupted document')})
        self.assertEqual(r.status_code,200,r.text);f=r.json();self.assertEqual(f['status'],'failed')
        self.assertEqual(self.client.get('/api/files/'+f['id']+'/download').content,b'corrupted document')

    def test_summary_versions_and_stale(self):
        self.case();self.chat();p=self.client.post('/api/projects',json={'title':'作品'}).json()
        self.client.post('/api/projects/move',json={'case_ids':[1],'project_id':p['id']})
        payload={'case_ids':[1],'file_ids':[],'goal':'汇总'}
        with patch.object(w.model_client,'complete',AsyncMock(return_value=({'content':'报告'},{}))):
            a=self.client.post('/api/projects/'+p['id']+'/summaries',json=payload).json()
            b=self.client.post('/api/projects/'+p['id']+'/summaries',json=payload).json()
        self.assertEqual([a['version'],b['version']],[1,2]);self.assertNotEqual(a['id'],b['id'])
        self.assertFalse(self.client.get('/api/projects/'+p['id']+'/summaries').json()['summaries'][0]['stale'])
        c=self.client.get('/api/cases/1').json();self.client.post('/api/cases',json={**c,'facts':{'payer':'CN'}})
        self.assertTrue(self.client.get('/api/projects/'+p['id']+'/summaries').json()['summaries'][0]['stale'])
        self.assertEqual(self.client.get('/api/summaries/'+a['id']+'/export').text.count(w.model_client.DISCLAIMER),1)

    def test_recommendation_off_independent_config_and_cache(self):
        self.case();c=self.chat();payload={'case_id':1,'message_id':c['active_leaf']}
        mock=AsyncMock(return_value=({'content':'{"questions":["需要哪些材料？","需要哪些材料？","还有哪些例外？"]}'},{}))
        with patch.object(w.model_client,'complete',mock):
            self.assertEqual(self.client.post('/api/recommendations/generate',json=payload).json()['questions'],[]);mock.assert_not_called()
            r=self.client.post('/api/recommendations/settings',json={'enabled':True,'provider':'deepseek','model':'rec-test','count':4,'api_key':'fake-rec-key'})
            self.assertNotIn('fake-rec-key',r.text)
            a=self.client.post('/api/recommendations/generate',json=payload);self.assertEqual(a.status_code,200,a.text)
            b=self.client.post('/api/recommendations/generate',json=payload)
            self.assertEqual(len(a.json()['questions']),2);self.assertTrue(b.json()['cached']);self.assertEqual(mock.await_count,1)
            self.assertEqual(mock.call_args.args[1][1],'rec-test')
        self.assertEqual(self.other.post('/api/recommendations/generate',json=payload).status_code,404)

    def test_web_search_uses_basic_and_preserves_snapshot(self):
        response=httpx_response({'results':[{'title':'官方','url':'https://www.iras.gov.sg/test','content':'text'}]})
        mock=AsyncMock(return_value=response)
        with patch.dict(w.os.environ,{'TAVILY_API_KEY':'fake-search-key'}),patch.object(w.httpx.AsyncClient,'post',mock):
            a=self.client.post('/api/web/search',json={'query':'测试'});self.assertEqual(a.status_code,200,a.text)
            b=self.client.post('/api/web/search',json={'query':'测试'});self.assertTrue(b.json()['cached'])
            self.assertEqual(mock.await_count,1);self.assertEqual(mock.call_args.kwargs['json']['search_depth'],'basic')
        self.assertEqual(self.client.get('/api/evidence/'+a.json()['evidence_id']).status_code,200)
        self.assertEqual(self.other.get('/api/evidence/'+a.json()['evidence_id']).status_code,404)

    def test_model_tool_result_is_real_backend_result(self):
        self.case();msg={'role':'assistant','content':None,'tool_calls':[{'id':'call1','type':'function','function':{'name':'legal_search','arguments':'{"query":"利息"}'}}]}
        model=AsyncMock(side_effect=[(msg,{}),({'role':'assistant','content':'依据检索结果'}, {})])
        search=AsyncMock(return_value={'query':'利息','results':[{'id':'actual-hit','excerpt':'真实返回文本'}],'evidence_id':'pinned'})
        with patch.object(w.model_client,'complete',model),patch.object(w,'legal_search',search):
            c=self.chat(prompt='查询利息税法')
        self.assertEqual(c['messages'][-1]['tools'][0]['output']['results'][0]['id'],'actual-hit')
        search.assert_awaited_once();self.assertEqual(c['messages'][-1]['evidence_ids'],['pinned'])

    def test_guest_migration_and_two_accounts(self):
        self.case();self.chat();self.other.get('/api/auth/session')
        async def identity(request):
            return {'id':request.headers['x-test-user'],'email':'test@example.invalid'} if request.headers.get('x-test-user') else None
        with patch.object(w,'auth_session',identity):
            self.assertEqual(self.client.get('/api/cases',headers={'x-test-user':'A'}).json()['cases'][0]['id'],1)
            self.assertEqual(self.other.get('/api/cases',headers={'x-test-user':'B'}).json()['cases'],[])
            self.assertEqual(self.other.get('/api/cases',headers={'x-test-user':'A'}).json()['cases'][0]['id'],1)

    def test_restart_keeps_completed_data(self):
        self.case();self.chat();f=self.client.post('/api/files',files={'file':('a.txt',b'persist')}).json()
        again=Store(path=self.db.path,dsn='')
        owner='guest:'+self.client.cookies['ct_session']
        self.assertEqual(len(again.get(owner,'case',1)['messages']),2)
        self.assertEqual(self.blob.get(again.get(owner,'file',f['id'])['object_key']),b'persist')

    def test_stop_preserves_partial_output(self):
        self.case();owner='guest:'+self.client.cookies['ct_session']
        async def slow(*args):
            yield {'type':'delta','text':'部分内容'}
            await asyncio.sleep(20)
        async def run():
            q=asyncio.Queue();task=asyncio.create_task(w.research_job(owner,1,{'prompt':'你好'},q))
            while True:
                e=await q.get()
                if e and e['type']=='delta':break
            task.cancel();await task
        with patch.object(w.model_client,'stream',slow):asyncio.run(run())
        c=self.db.get(owner,'case',1);self.assertEqual(c['messages'][-1]['status'],'interrupted')
        self.assertIn('部分内容',c['messages'][-1]['text'])

    def test_origin_rejected(self):
        self.assertEqual(self.client.post('/api/cases',json={},headers={'Origin':'https://evil.invalid'}).status_code,403)


def httpx_response(data):
    import httpx
    return httpx.Response(200,json=data)


class CalculationTests(unittest.TestCase):
    def test_decimal_explicit_assumption_and_unpublished_gate(self):
        facts={'payer':'CN','recipient':'SG','income_type':'ROYALTIES','transaction_subtype':'equipment','date':'2026-10-07','classification_confirmed':True,'beneficial_owner':True,'recipient_tax_resident':True,'eligibility_documents':True,'principal_purpose_test_passed':True}
        with patch.object(legal_service,'rows',return_value=[]):
            r=legal_service.assess({'facts':facts,'amount':'0.10','assumptions':{'nominal_rate':'0.15'}})
        self.assertIsNone(r['tax_estimate']);self.assertEqual(r['conditional_example']['amount'],'0.02')
        self.assertEqual(r['status'],'rule_not_published')

    def test_pe_blocks_simple_withholding_formula(self):
        r=legal_service.assess({'facts':{'pe_effective_connection':True},'amount':'100','assumptions':{'nominal_rate':'0.1'}})
        self.assertEqual(r['status'],'exception_requires_analysis');self.assertIsNone(r['conditional_example'])

    def test_missing_classification_is_explicit(self):
        r=legal_service.assess({'facts':{},'amount':'100'})
        self.assertIn('classification_confirmed',r['missing_facts'])
        self.assertIsNone(r['tax_estimate'])

    def test_invalid_amount_rejected(self):
        for value in ['NaN','Infinity','-1']:
            with self.assertRaises(ValueError):legal_service.assess({'amount':value})


if __name__=='__main__':unittest.main()
